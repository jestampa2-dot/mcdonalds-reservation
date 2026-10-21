from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\User\mcdonalds-reservation")
SOURCE = ROOT / "Full Thesis - Final RRL 20 Local 20 Foreign Hyperlinked.docx"
OUTPUT = ROOT / "Full Thesis - Final RRL 20 Local 20 Foreign Hyperlinked - Updated 1.3.docx"


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def append_after(anchor, text="", style_name="Normal", first_indent=True):
    from docx.text.paragraph import Paragraph

    new_element = OxmlElement("w:p")
    anchor._p.addnext(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    paragraph.style = style_name
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if style_name == "Normal" else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(6)
    if style_name == "Normal" and first_indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    return paragraph


def build():
    doc = Document(str(SOURCE))
    paragraphs = list(doc.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "1.3 General Objectives of the Study")
    end = next(i for i, p in enumerate(paragraphs[start + 1 :], start + 1) if p.text.strip().startswith("1.4 Scope"))
    heading = paragraphs[start]

    for paragraph in paragraphs[start + 1 : end]:
        delete_paragraph(paragraph)

    anchor = heading
    anchor = append_after(
        anchor,
        "This study aims to develop a digital event reservation system for McDonald's that will improve the management of birthday parties, business meetings, and other event reservations. Specifically, the study aims to accomplish the following objectives:",
    )
    objectives = [
        "1. To gather data and determine the system requirements needed for the proposed McDonald's Event Planner.",
        "2. To design and develop a digital event reservation system for managing bookings, schedules, packages, and customer information.",
        "3. To test and evaluate the system in terms of usability, reliability, efficiency, and its ability to support event reservation management.",
    ]
    for objective in objectives:
        p = append_after(anchor, objective, first_indent=False)
        p.paragraph_format.left_indent = Inches(0.5)
        anchor = p

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
